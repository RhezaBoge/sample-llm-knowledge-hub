"""Step 1 of the pipeline: get plain text out of an uploaded file."""

import io
from pathlib import Path

SUPPORTED = (".txt", ".md", ".pdf", ".docx")


def read_document(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()

    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    raise ValueError(f"Unsupported file type '{ext}'. Please use one of: {', '.join(SUPPORTED)}")
